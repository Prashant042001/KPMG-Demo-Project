"""Regression tests for the reported inconsistent derivations and the optional-supporting-file workflow."""
from __future__ import annotations

import pytest

from app.derivation.engine import CANNOT_BE_DERIVED
from samples import synthetic as syn

REPORTED = [
    ("9925MTN", "T9925", "9925", "R3"),
    ("4908MTN", "T4908", "4908", "R3"),
    ("3133MTN", "T3133", "3133", "R3"),
    ("9986MTN/S", "T9986", "9986", "R10"),
    ("21353MTN/E", "T21353", "21353", "R11"),
    ("MTN6364", "T6364", "6364", "R12"),
    ("T8736", "T8736", "8736", "R2"),
    ("T8646/ELEC", "T8646", "8646", "R5"),
    ("T7386/E", "T7386", "7386", "R6"),
]


@pytest.mark.parametrize("code,t,site,rule", REPORTED, ids=[r[0] for r in REPORTED])
def test_reported_codes_without_any_supporting_file(run, code, t, site, rule):
    res = run(source=syn.base_workbook([["I", code, None, 1, "x"]]), bts=False, ops=False)
    o = res.result.records[0].outcome
    assert (o.tnumber, o.site_id, o.t_rule) == (t, site, rule)


@pytest.mark.parametrize("code,t,site,rule", REPORTED, ids=[r[0] for r in REPORTED])
def test_reported_codes_are_not_blocked_by_failed_reference_lookups(derive, code, t, site, rule):
    """BTS row present for the code but with a blank / invalid T-number must not stop the pattern rule."""
    _, outs = derive([(code, "")], bts_rows=[[code, "", "", "a"], [code, "not-a-t", "", "b"]])
    assert (outs[0].tnumber, outs[0].site_id, outs[0].t_rule) == (t, site, rule)


def test_equivalent_numeric_prefix_codes_use_the_same_rule(derive):
    codes = ["9925MTN", "4908MTN", "1000MTN", "12345MTN"]
    _, outs = derive([(c, "") for c in codes])
    assert [(o.tnumber, o.t_rule) for o in outs] == [("T9925", "R3"), ("T4908", "R3"), ("T1000", "R3"), ("T12345", "R3")]


@pytest.mark.parametrize(
    "raw",
    ["﻿4908MTN", "4908MTN​", "4908 MTN", " 4908 MTN ", "４９０８MTN", "4908mtn", "⁠4908MTN­"],
)
def test_hidden_characters_and_spacing_do_not_break_derivation(derive, raw):
    res, outs = derive([(raw, "")])
    assert (outs[0].tnumber, outs[0].site_id, outs[0].t_rule) == ("T4908", "4908", "R3")
    assert res.result.records[0].vendor_site_code == raw.strip()  # original preserved (only outer trim for display)


def test_numeric_excel_cell_values(derive):
    _, outs = derive([(4908, ""), (40123.0, "")])
    # A bare number is not one of the approved patterns: never fabricated.
    assert [o.tnumber for o in outs] == [CANNOT_BE_DERIVED, CANNOT_BE_DERIVED]


def test_explicit_t_number_outranks_composed_reference_value(derive):
    """R28 only supplies a Site ID; 'T' + that Site ID must not overwrite R3's explicit T-number."""
    res, outs = derive([("4908MTN", "")], ops_rows=[["4908MTN", "4909", "a"]])
    o = outs[0]
    assert (o.tnumber, o.site_id, o.t_rule, o.site_rule) == ("T4908", "4908", "R3", "R3")  # consistent pair
    assert o.candidate_tnumbers == "T4908, T4909"
    assert any("R28 gave 4909" in e.detail for e in res.result.exceptions)  # disagreement recorded


def test_r7_r21_remain_flagged_until_suffixes_are_approved(run):
    res = run()
    for rid in ("R7", "R21"):
        r = res.ruleset.by_id(rid)
        assert r.status == "REQUIRES_CONFIRMATION" and "/ELEC" in r.status_reason
    assert res.ruleset.by_id("R7").source_column == "VENDOR_SITE_CODE"
    assert res.ruleset.by_id("R21").source_column == "VENDOR_SITE_CODE_ALT"


def test_r7_executes_once_suffixes_are_approved(derive, config_with):
    cfg = config_with(lambda c: c["recognised_suffixes"].update(values=["/S"]))
    res, outs = derive([("T1234/S", "")], config=cfg)
    assert res.ruleset.by_id("R7").status == "ACTIVE"
    assert {c.rule_id for c, _ in outs[0].candidates} >= {"R7"}
    assert outs[0].tnumber == "T1234"


def test_r7_does_not_strip_unapproved_text(derive, config_with):
    cfg = config_with(lambda c: c["recognised_suffixes"].update(values=["/S"]))
    _, outs = derive([("T1234/X", "")], config=cfg)
    assert "R7" not in {c.rule_id for c, _ in outs[0].candidates}


def test_source_only_run_reports_missing_reference_per_record(run):
    res = run(source=syn.base_workbook([["I-1", "SYN-001", None, 1, "x"], ["I-2", "T6542", None, 2, "x"]]), bts=False, ops=False)
    first, second = res.result.records
    assert first.outcome.tnumber == CANNOT_BE_DERIVED  # only BTS could have resolved it
    assert "BTS Site Report Schedule" in first.exception_reason
    assert second.outcome.tnumber == "T6542" and "not supplied" not in second.exception_reason
    cats = {e.category for e in res.result.exceptions}
    assert "MISSING_REFERENCE_DATA" in cats


def test_missing_source_value_is_its_own_category(derive):
    res, _ = derive([("", "")])
    assert "MISSING_SOURCE_VALUE" in {e.category for e in res.result.exceptions}


def test_derivation_counts_by_rule_are_measured(run):
    m = run().result.metrics
    assert sum(m["tnumber_derived_by_rule"].values()) == m["tnumber_derived"]
    assert m["tnumber_match_rate_pct"] == round(100 * m["tnumber_derived"] / m["input_rows"], 2)


def test_diagnostic_columns_present(run):
    import io

    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(run().outputs["master"].content))["Derivation Trace"]
    header = [c.value for c in ws[1]]
    for col in ("Original Vendor Site Code", "Original VENDOR_SITE_CODE_ALT", "Candidate T-number(s)", "TNUMBER", "SITE ID",
                "T-number Rule ID", "T-number Derivation Source", "T-number Derivation Status", "Conflict Indicator", "Exception Reason"):
        assert col in header


VENDOR_SITE_HEADER = ["Vendor Name", "Vendor Site", "Invoice Number", "Invoice Amount"]


def test_payment_extract_with_vendor_site_column_applies_pattern_rules(run):
    """Root cause of the 41/68 result: the source had 'Vendor Site' but no VENDOR_SITE_CODE, so R2-R14 never ran."""
    rows = [["Synthetic Co", code, f"PN-{i}", 100] for i, code in enumerate(["4908MTN", "7572MTN", "T21151/E", "T8736", "MTN01/E T3113"])]
    res = run(source=syn.base_workbook(rows, header=VENDOR_SITE_HEADER), bts=False, ops=False)
    got = [(r.outcome.tnumber, r.outcome.t_rule) for r in res.result.records]
    assert got == [("T4908", "R3"), ("T7572", "R3"), ("T21151", "R6"), ("T8736", "R2"), ("T3113", "R8")]
    assert any("read 'Vendor Site'" in n for n in res.summary["notes"])


def test_vendor_site_fallback_can_be_disabled(run, config_with):
    cfg = config_with(lambda c: c.update(source_column_fallbacks={}))
    rows = [["Synthetic Co", "4908MTN", "PN-1", 100]]
    res = run(source=syn.base_workbook(rows, header=VENDOR_SITE_HEADER), bts=False, ops=False, config=cfg)
    assert res.result.records[0].outcome.tnumber == CANNOT_BE_DERIVED


SUPPLIER_HEADER = ["VENDOR_SITE_CODE", "VENDOR_SITE_CODE_ALT", "T_NUMBER"]


def test_x_code_in_alt_with_trailing_text_is_found(run):
    """R4/R17: 'begins with X followed by a numeric identifier' - trailing text must not stop the match."""
    supplier = syn.bts_workbook([["4000000001", "X100999 Synthetic Hill TX", "NA"]], header=SUPPLIER_HEADER)
    rows = [["Co", "4000000001", "PN-1", 1]]
    res = run(source=syn.base_workbook(rows, header=VENDOR_SITE_HEADER), bts=supplier, ops=False)
    o = res.result.records[0].outcome
    r17 = [c for c, _ in o.candidates if c.rule_id == "R17"]
    assert r17 and r17[0].matching_key == "X100999"
    assert (o.tnumber, o.site_id, o.t_rule) == ("X100999", "X100999", "R17")  # mapped per x_codes setting


def test_x_code_in_alt_validated_by_reference_becomes_site_id(run):
    supplier = syn.bts_workbook([["4000000001", "X100999 Synthetic Hill TX", "NA"]], header=SUPPLIER_HEADER)
    ops = syn.ops_workbook([["ACC-9", "X100999", "Synthetic"]])
    res = run(source=syn.base_workbook([["Co", "4000000001", "PN-1", 1]], header=VENDOR_SITE_HEADER), bts=supplier, ops=ops)
    o = res.result.records[0].outcome
    assert (o.site_id, o.site_rule, o.tnumber) == ("X100999", "R17", "X100999")
    assert not any("without reference validation" in e.detail for e in res.result.exceptions)


def test_code_column_x_code_takes_precedence_over_alt(run):
    """Business order: VENDOR_SITE_CODE first, VENDOR_SITE_CODE_ALT only if nothing found there."""
    rows = [["I-1", "X200001", "X300003 OTHER", 1, "x"], ["I-2", "T4321", "X300003", 1, "x"], ["I-3", "PLAIN", "X300003", 1, "x"]]
    res = run(source=syn.base_workbook(rows), bts=False, ops=False)
    got = [(r.outcome.tnumber, r.outcome.t_rule) for r in res.result.records]
    assert got == [("X200001", "R4"), ("T4321", "R2"), ("X300003", "R17")]


def test_x_codes_without_reference_not_mapped_when_disabled(run, config_with):
    cfg = config_with(lambda c: c["x_codes"].update(accept_without_reference=False))
    supplier = syn.bts_workbook([["4000000001", "X100999 Synthetic Hill TX", "NA"]], header=SUPPLIER_HEADER)
    res = run(source=syn.base_workbook([["Co", "4000000001", "PN-1", 1]], header=VENDOR_SITE_HEADER), bts=supplier, ops=False, config=cfg)
    o = res.result.records[0].outcome
    assert o.tnumber == CANNOT_BE_DERIVED and o.site_status == "UNVERIFIED"


@pytest.mark.parametrize("marker", ["NA", "n/a", "#N/A", "NULL", "-"])
def test_reference_placeholders_are_blank_not_invalid(run, marker):
    supplier = syn.bts_workbook([["4908MTN", "", marker]], header=SUPPLIER_HEADER)
    res = run(source=syn.base_workbook([["Co", "4908MTN", "PN-1", 1]], header=VENDOR_SITE_HEADER), bts=supplier, ops=False)
    assert res.result.records[0].outcome.tnumber == "T4908"
    assert "INVALID_CANDIDATE" not in {e.category for e in res.result.exceptions}


def test_site_id_stays_paired_with_selected_t_number(run):
    """Code says 12470, ALT says T12740: the code rule wins by priority and the pair stays consistent."""
    supplier = syn.bts_workbook([["12470MTN SYN", "T12740 SYNTHETIC MALL", ""]], header=SUPPLIER_HEADER)
    res = run(source=syn.base_workbook([["Co", "12470MTN SYN", "PN-1", 1]], header=VENDOR_SITE_HEADER), bts=supplier, ops=False)
    o = res.result.records[0].outcome
    assert (o.tnumber, o.site_id, o.t_rule, o.site_rule) == ("T12470", "12470", "R13", "R13")
    assert not o.conflict and o.candidate_tnumbers == "T12470, T12740"


@pytest.mark.parametrize(
    "code,t,site",
    [("T000-1234", "T1234", "1234"), ("T000/1234", "T1234", "1234"), ("T0-56789", "T56789", "56789"),
     ("T0000/4321", "T4321", "4321"), ("t000 - 4321", "T4321", "4321")],
)
def test_r30_zero_prefix_with_delimiter(derive, code, t, site):
    _, outs = derive([(code, "")])
    assert (outs[0].tnumber, outs[0].site_id, outs[0].t_rule) == (t, site, "R30")
    assert not outs[0].conflict


@pytest.mark.parametrize("code,t,site", [("T590MTN", "T590", "590"), ("T590 MTN", "T590", "590"), ("T59012MTN", "T59012", "59012")])
def test_r31_t_number_followed_by_mtn(derive, code, t, site):
    _, outs = derive([(code, "")])
    assert (outs[0].tnumber, outs[0].site_id) == (t, site)
    assert "R31" in {c.rule_id for c, _ in outs[0].candidates}


@pytest.mark.parametrize("code", ["T000-1234-X", "T000-123", "T000-1234567", "AT000-1234", "T590MTN/S", "T59MTN", "T590MTNX", "T590"])
def test_new_rules_only_match_their_exact_shapes(derive, code):
    _, outs = derive([(code, "")])
    assert not {"R30", "R31"} & {c.rule_id for c, _ in outs[0].candidates}


@pytest.mark.parametrize("code", ["T0000", "T0000-X", "X000-1234", "ABC-T00000"])
def test_all_zero_placeholders_are_never_identifiers(derive, code):
    _, outs = derive([(code, "")])
    assert outs[0].tnumber == CANNOT_BE_DERIVED and outs[0].site_id == ""


def test_new_rules_read_vendor_site_column_too(run):
    rows = [["Co", "T000-1234", "PN-1", 1], ["Co", "T590MTN", "PN-2", 1]]
    res = run(source=syn.base_workbook(rows, header=VENDOR_SITE_HEADER), bts=False, ops=False)
    assert [r.outcome.tnumber for r in res.result.records] == ["T1234", "T590"]

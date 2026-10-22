"""REST API end-to-end tests (in-process, synthetic data)."""
from __future__ import annotations

import io
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import Settings, get_settings
from app.main import app
from samples import synthetic as syn

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def client():
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def _files(supporting=("bts", "ops")):
    files = [("source_file", ("any_source_name.xlsx", syn.base_workbook(), XLSX))]
    builders = {"bts": syn.bts_workbook, "ops": syn.ops_workbook}
    for i, kind in enumerate(supporting):
        files.append(("supporting_files", (f"ref_{i}.xlsx", builders[kind](), XLSX)))
    return files


def _wait(client, job_id, timeout=30):
    deadline = time.time() + timeout
    while time.time() < deadline:
        body = client.get(f"/api/jobs/{job_id}").json()
        if body["status"] in ("completed", "failed"):
            return body
        time.sleep(0.1)
    raise AssertionError("job did not finish")


def test_health(client):
    r = client.get("/api/health")
    assert r.status_code == 200 and r.json()["status"] == "ok"


def test_config_exposes_backend_rules_and_two_upload_kinds(client):
    body = client.get("/api/config").json()
    assert body["scope"] == ["TNUMBER", "SITE ID"]
    assert set(body["uploads"]) == {"source_file", "supporting_files"}
    assert body["rules"]["source"].startswith("backend configuration")
    assert "bts" in body["datasets"] and "rules" not in body["datasets"]


def test_validate_ok(client):
    body = client.post("/api/validate", files=_files()).json()
    assert body["valid"], body
    assert body["source_rows"] == len(syn.BASE_ROWS)
    assert {"R7", "R21"} <= {x["rule_id"] for x in body["rules"]["requires_confirmation"]}


def test_validate_source_only_is_valid(client):
    body = client.post("/api/validate", files=_files(supporting=())).json()
    assert body["valid"] and body["errors"] == []
    assert "BTS Site Report Schedule" in body["missing_supporting_datasets"]


def test_validate_requires_source_file(client):
    files = [("supporting_files", ("ref.xlsx", syn.bts_workbook(), XLSX))]
    body = client.post("/api/validate", files=files).json()
    assert not body["valid"] and body["errors"] == ["A source file is required."]


@pytest.mark.parametrize("supporting", [(), ("bts",), ("bts", "ops")], ids=["none", "one", "several"])
def test_process_with_zero_one_or_several_supporting_files(client, supporting):
    r = client.post("/api/process", files=_files(supporting), data={"strict_matching": "false", "output_format": "xlsx"})
    assert r.status_code == 202
    job = _wait(client, r.json()["job_id"])
    assert job["status"] == "completed", job
    assert job["metrics"]["input_rows"] == len(syn.BASE_ROWS)
    by_code = {p["vendor_site_code"]: p["tnumber"] for p in job["preview"]}
    assert by_code["4908MTN"] == "T4908" and by_code["9925MTN"] == "T9925"
    expected_syn001 = "T9001" if "bts" in supporting else "cannot be derived"
    assert by_code["SYN-001"] == expected_syn001


def test_process_and_download(client):
    r = client.post("/api/process", files=_files(), data={"output_format": "csv"})
    job = _wait(client, r.json()["job_id"])
    assert job["status"] == "completed"
    for kind in ("master", "mapping", "exceptions"):
        d = client.get(f"/api/jobs/{job['job_id']}/download/{kind}")
        assert d.status_code == 200 and "attachment" in d.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(client.get(f"/api/jobs/{job['job_id']}/download/master").content))
    assert wb.sheetnames == ["Master Database", "Mapping Summary", "Exceptions", "Processing Summary", "Rule Audit"]


def test_legacy_rules_file_field_is_ignored(client):
    files = _files() + [("rules_file", ("T_Number_Rules.xlsx", b"PK\x03\x04garbage", XLSX))]
    job = _wait(client, client.post("/api/process", files=files).json()["job_id"])
    assert job["status"] == "completed"


def test_process_failure_is_reported_without_success(client):
    files = [("source_file", ("bad.xlsx", syn.base_workbook([["I", "x"]], header=["A", "B"]), XLSX))]
    job = _wait(client, client.post("/api/process", files=files).json()["job_id"])
    assert job["status"] == "failed" and job["validation"]["errors"]
    assert client.get(f"/api/jobs/{job['job_id']}/download/master").status_code == 409


def test_unknown_and_malformed_job_ids(client):
    assert client.get("/api/jobs/" + "a" * 32).status_code == 404
    assert client.get("/api/jobs/../../etc").status_code in (404, 422)
    assert client.get("/api/jobs/not-a-job").status_code == 422


def test_rejects_unsupported_extension(client):
    r = client.post("/api/validate", files=[("source_file", ("evil.exe", b"MZ", "application/octet-stream"))])
    assert r.status_code == 422 and "unsupported file type" in r.json()["message"]


def test_file_size_limit(client):
    app.dependency_overrides[get_settings] = lambda: Settings(max_file_size_mb=1)
    big = b"a,b\n" + b"x,y\n" * 300_000
    r = client.post("/api/validate", files=[("source_file", ("big.csv", big, "text/csv"))])
    assert r.status_code == 413


def test_errors_do_not_expose_paths(client):
    r = client.post("/api/validate", files=[("source_file", ("bad.xlsx", b"PK\x03\x04garbage", XLSX))])
    assert "site-packages" not in r.text and ":\\\\" not in r.text


def test_root_points_to_health_and_docs(client):
    assert client.get("/").json()["health"] == "/api/health"

"""Output workbook structure, preservation of source data and formula-injection protection."""
from __future__ import annotations

import datetime as dt
import io

from openpyxl import load_workbook

from app.output.master_layout import MASTER_HEADERS
from app.services.gis import payment_type
from app.output.writer import TRACE_COLUMNS, to_csv
from samples import synthetic as syn

SHEETS = ["Master Database", "Derivation Trace", "Mapping Summary", "Exceptions", "Processing Summary", "Rule Audit"]
as_excel = lambda v: dt.datetime(v.year, v.month, v.day) if type(v) is dt.date else v  # Excel stores dates as datetimes


def _master(res):
    return load_workbook(io.BytesIO(res.outputs["master"].content))


def _rows(ws):
    header = [c.value for c in ws[1]]
    return [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)]


def test_output_worksheet_names(run):
    assert _master(run()).sheetnames == SHEETS


def test_master_has_exactly_the_final_columns_in_order(run):
    ws = _master(run())["Master Database"]
    header = [c.value for c in ws[1]]
    assert ws.freeze_panes == "A2"
    assert header == MASTER_HEADERS and len(header) == 64
    assert header[:4] == ["TNUMBER", "SITE ID", "Vendor Name", "Vendor Site Code"]
    assert header[-5:] == ["Site Status", "Site Name", "Site Owner", "Latitude", "Longitude"]


def test_base_payment_data_is_copied_unchanged(run):
    res = run()
    rows = _rows(_master(res)["Master Database"])
    source = syn.base_rows_for_sheet()
    assert len(rows) == len(source) == res.summary["results"]["output_rows"]
    for out, src in zip(rows, source):
        # synthetic header: INVOICE_NUM, VENDOR_SITE_CODE, ALT, INVOICE_AMOUNT, DESCRIPTION, Vendor Name, Payment Date, Payment Amount
        assert out["Invoice Number"] == src[0] and out["Vendor Site Code"] == src[1] and out["Invoice Amount"] == src[3]
        assert out["Vendor Name"] == src[5] and out["Payment Date"] == as_excel(src[6]) and out["Payment Amount"] == src[7]
        assert out["Check Number"] is None  # not in the synthetic source file
    assert "Check Number" in " ".join(res.summary["notes"])


def test_derived_values_in_master_with_na_for_underived_tnumber(run):
    rows = _rows(_master(run())["Master Database"])
    expected = [("NA" if t == "cannot be derived" else t, s) for _, _, t, s in syn.BASE_ROWS]
    assert [(r["TNUMBER"], r["SITE ID"] or "") for r in rows] == expected
    assert any(t == "NA" for t, _ in expected)


def test_operational_columns_come_from_the_site_list(run):
    header = ["Account Number", "Site ID", "Site Name", "Site Status", "Site Owner", "Latitude", "Longitude"]
    ops = syn.ops_workbook([["OPS-1", "6542", "Synthetic Site G", "On Air", "Synthetic Owner G", -29.8587, 31.0218]], header=header)
    rows = {r["Vendor Site Code"]: r for r in _rows(_master(run(ops=ops))["Master Database"])}
    hit = rows["T6542"]
    assert (hit["Site Status"], hit["Site Name"], hit["Site Owner"], hit["Latitude"], hit["Longitude"]) == \
        ("On Air", "Synthetic Site G", "Synthetic Owner G", -29.8587, 31.0218)
    miss = rows["3133MTN"]  # SITE ID 3133 is not in this site list
    assert (miss["Site Status"], miss["Site Name"], miss["Latitude"]) == ("NA", "NA", "NA")


def test_randomized_fields_are_filled_and_consistent_per_vendor(run):
    rows = _rows(_master(run())["Master Database"])
    for r in rows:
        assert r["Bank Name"] and r["Bank Account Num"] and r["Vat Registration Num"] and r["E-mail ID"]
        assert r["Bank Account Name"] == r["Vendor Name"]
        assert r["Frequency Code"] in ("Monthly", "Annually") and r["Country"] == "South Africa"
        assert len(r["ZIP (4 Digit Only)"]) == 4 and r["ZIP (4 Digit Only)"].isdigit()
        assert r["Lease Execution date"] < r["Lease start date"] < r["Termination Date"]
        assert r["START_DATE_ACTIVE"] < r["END_DATE_ACTIVE"]
    by_vendor = {}
    for r in rows:
        by_vendor.setdefault(r["Vendor Name"], set()).add((r["Bank Name"], r["Bank Account Num"], r["Vendor ID"]))
    assert all(len(v) == 1 for v in by_vendor.values()) and len(by_vendor) > 1
    again = _rows(_master(run())["Master Database"])
    assert [r["Bank Account Num"] for r in again] == [r["Bank Account Num"] for r in rows]  # stable across runs


def test_payment_type_from_invoice_number():
    assert payment_type("PN-1001") == payment_type(" pn778") == "Rental + Fixed"
    assert payment_type("INV-1") == payment_type(None) == "Variable Electricity"


def test_derivation_trace_sheet_keeps_traceability(run):
    ws = _master(run())["Derivation Trace"]
    assert [c.value for c in ws[1]] == TRACE_COLUMNS
    rows = _rows(ws)
    assert [r["TNUMBER"] for r in rows] == [t for _, _, t, _ in syn.BASE_ROWS]  # internal value, not NA


def test_formula_text_is_stored_as_text_not_formula(run):
    src = syn.base_workbook([["PN-1", "T6542", None, 1, "x", '=HYPERLINK("http://example.invalid","x")', None, 5]])
    ws = _master(run(source=src))["Master Database"]
    cell = ws.cell(2, MASTER_HEADERS.index("Vendor Name") + 1)
    assert cell.value.startswith("=HYPERLINK") and cell.data_type == "s"


def test_csv_reports_neutralise_formulas():
    out = to_csv(["a", "b", "c"], [["=1+1", "-12.5", "@cmd"]]).decode("utf-8-sig")
    assert "'=1+1" in out and ",-12.5," in out and "'@cmd" in out


def test_csv_report_format(run):
    res = run(fmt="csv")
    assert res.outputs["mapping"].filename.endswith(".csv")
    assert res.outputs["exceptions"].content.decode("utf-8-sig").startswith("Category,Severity")
    assert res.outputs["master"].filename.endswith(".xlsx")


def test_rule_audit_and_summary_content(run):
    wb = _master(run())
    audit = {r[0]: r for r in wb["Rule Audit"].iter_rows(min_row=2, values_only=True)}
    assert audit["R7"][12] == "Requires business confirmation"
    assert audit["R2"][8] >= 1  # successful matches
    metrics = {r[0]: r[1] for r in wb["Processing Summary"].iter_rows(min_row=2, values_only=True) if r[0]}
    assert metrics["input_rows"] == len(syn.BASE_ROWS)
    assert "tnumber_match_rate_pct" in metrics


def test_exceptions_sheet_lists_unresolved_records(run):
    wb = _master(run())
    cats = {r[0] for r in wb["Exceptions"].iter_rows(min_row=2, values_only=True)}
    assert {
        "UNRESOLVED_TNUMBER", "UNRESOLVED_SITE_ID", "AMBIGUOUS_REFERENCE_MAPPING", "MULTIPLE_CANDIDATE_IDENTIFIERS",
        "RESOLVED_BY_PRIORITY", "UNVERIFIED_IDENTIFIER", "RULE_REQUIRES_CONFIRMATION",
    } <= cats



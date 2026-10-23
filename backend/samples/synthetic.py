"""Synthetic, non-sensitive sample data for tests and local demonstrations.

Every vendor site code, T-number, account number and amount below is invented.
Run ``python -m samples.synthetic`` from the ``backend`` folder to write the sample files to ``samples/output``.
"""
from __future__ import annotations

import datetime as dt
import io
from pathlib import Path

from openpyxl import Workbook

VSC, ALT = "VENDOR_SITE_CODE", "VENDOR_SITE_CODE_ALT"

BASE_HEADER = ["INVOICE_NUM", VSC, ALT, "INVOICE_AMOUNT", "DESCRIPTION", "Vendor Name", "Payment Date", "Payment Amount"]
# (vendor site code, alt, expected TNUMBER, expected SITE ID) - the expectation columns are for tests only.
BASE_ROWS = [
    ("SYN-001", "", "T9001", "9001"),            # R1 reference mapping
    ("T6542", "", "T6542", "6542"),              # R2
    ("3133MTN", "", "T3133", "3133"),            # R3
    ("X8754", "", "X8754", "X8754"),             # R4 validated against Operational Site List
    ("X1111", "", "X1111", "X1111"),             # R4 not in reference data: mapped per x_codes setting, flagged Info
    ("T8646/ELEC", "", "T8646", "8646"),         # R5
    ("T7386/E", "", "T7386", "7386"),            # R6
    ("ABC-T8102-RENT", "", "T8102", "8102"),     # R8
    ("T12345", "", "T12345", "12345"),           # R2 + R9 (5 digits)
    ("4455MTN/S", "", "T4455", "4455"),          # R10 (T composed)
    ("4456MTN/E", "", "T4456", "4456"),          # R11
    ("MTN77881", "", "T77881", "77881"),         # R12
    ("5566-LEASE", "", "T5566", "5566"),         # R13
    ("0123", "", "cannot be derived", "123"),    # R14 (T123 fails T-number format)
    ("UNKNOWN-CODE", "T2468", "T2468", "2468"),  # R15 via ALT
    ("NOPE", "", "cannot be derived", ""),       # nothing matches
    ("T1111-T2222", "", "cannot be derived", ""),  # R8 multiple T-numbers -> ambiguous
    ("SYN-CONF", "", "cannot be derived", ""),   # R1 conflicting reference rows -> ambiguous
    ("ACC-500", "", "T6600", "6600"),            # R28 Operational Site List
    ("T5000", "", "T5001", "5001"),              # R1 (T5001) beats R2 (T5000) -> conflict flagged
    ("SYN-001", "", "T9001", "9001"),            # duplicate source record is preserved
    # Codes reported from a real run (as patterns only; no real data):
    ("9925MTN", "", "T9925", "9925"),            # R3
    ("4908MTN", "", "T4908", "4908"),            # R3 - same rule as 9925MTN
    ("9986MTN/S", "", "T9986", "9986"),          # R10
    ("21353MTN/E", "", "T21353", "21353"),       # R11, 5 digits
    ("MTN6364", "", "T6364", "6364"),            # R12
    ("T8736", "", "T8736", "8736"),              # R2
    ("﻿4908MTN​", "", "T4908", "4908"),  # invisible characters from an export
    (" 4908 MTN ", "", "T4908", "4908"),         # stray spaces
]

BTS_HEADER = [VSC, "T number", ALT, "Site Name"]
BTS_ROWS = [
    ["SYN-001", "T9001", "", "Synthetic Site A"],
    ["SYN-001", "T9001", "", "Synthetic Site A"],  # duplicate, agreeing
    ["SYN-CONF", "T3001", "", "Synthetic Site B"],
    ["SYN-CONF", "T3002", "", "Synthetic Site C"],  # conflicting
    ["T5000", "T5001", "", "Synthetic Site D"],
]

OPS_HEADER = ["Account Number", "Site ID", "Site Name", "Region", "Property Owner", "Latitude", "Longitude"]
# Invented sites and coordinates. 'OPS-' accounts match no source code, so they only feed the GIS view.
OPS_ROWS = [
    ["ACC-500", "6600", "Synthetic Site E", "Synthetic North", "Synthetic Owner E", -26.1950, 28.0340],
    ["ACC-777", "X8754", "Synthetic Site F", "Synthetic North", "Synthetic Owner F", -25.7460, 28.1880],
    ["OPS-9001", "9001", "Synthetic Site A", "Synthetic North", "Synthetic Owner A", -26.1076, 28.0567],
    ["OPS-6542", "6542", "Synthetic Site G", "Synthetic Coast", "Synthetic Owner G", -29.8587, 31.0218],
    ["OPS-3133", "3133", "Synthetic Site H", "Synthetic Cape", "Synthetic Owner H", -33.9249, 18.4241],
    ["OPS-8646", "8646", "Synthetic Site I", "Synthetic Cape", "Synthetic Owner I", "-33,9600", "18,4700"],
    ["OPS-7777", "7777", "Synthetic Site J (no payments)", "Synthetic Central", "Synthetic Owner J", -29.0852, 26.1596],
    ["OPS-5555", "5555", "Synthetic Site K (no coordinates)", "Synthetic Central", "Synthetic Owner K", None, None],
]


def _xlsx(sheets: dict[str, list[list]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row_no, r in enumerate(rows, start=1):
            ws.append(r)
            for col_no, value in enumerate(r, start=1):
                if isinstance(value, str) and value.startswith("="):
                    ws.cell(row_no, col_no).data_type = "s"  # literal text, as typed with a leading apostrophe
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def base_rows_for_sheet() -> list[list]:
    rows = []
    for i, (vsc, alt, _, _) in enumerate(BASE_ROWS, start=1):
        desc = '=HYPERLINK("http://example.invalid","x")' if i == 1 else f"Synthetic payment {i}"
        rows.append([f"INV-{i:04d}", vsc, alt or None, 1000 + i, desc,
                     f"Synthetic Vendor {i % 5 + 1}", dt.date(2025, 1 + i % 12, 1 + i % 28), 1000.0 + 10 * i])
    return rows


def base_workbook(rows: list[list] | None = None, header: list[str] | None = None) -> bytes:
    # Title rows above the header exercise header-row detection.
    body = rows if rows is not None else base_rows_for_sheet()
    return _xlsx({"Payments": [["Synthetic Payment Report"], ["Generated for testing"], header or BASE_HEADER, *body]})


def bts_workbook(rows: list[list] | None = None, header: list[str] | None = None) -> bytes:
    return _xlsx({"BTS Site Report": [header or BTS_HEADER, *(rows if rows is not None else BTS_ROWS)]})


def ops_workbook(rows: list[list] | None = None, header: list[str] | None = None) -> bytes:
    return _xlsx({"Summary": [["Operational Site List Summary"], header or OPS_HEADER, *(rows if rows is not None else OPS_ROWS)]})


def write_samples(target: Path) -> list[Path]:
    target.mkdir(parents=True, exist_ok=True)
    files = {
        "Synthetic_Source_File.xlsx": base_workbook(),
        "Synthetic_BTS_Site_Report_Schedule.xlsx": bts_workbook(),
        "Synthetic_Operational_Site_List_Summary.xlsx": ops_workbook(),
    }
    paths = []
    for name, content in files.items():
        p = target / name
        p.write_bytes(content)
        paths.append(p)
    return paths


if __name__ == "__main__":
    out = Path(__file__).resolve().parent / "output"
    for p in write_samples(out):
        print(p)
